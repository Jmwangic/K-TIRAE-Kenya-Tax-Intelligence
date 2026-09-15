from docx import Document
from docx.shared import Pt

output_path = r'd:\KRA W\dashboard_breakdown.docx'

doc = Document()
section = doc.sections[0]
section.left_margin = 720
section.right_margin = 720
section.top_margin = 720
section.bottom_margin = 720

p = doc.add_paragraph()
r = p.add_run('KRA Tax Anomaly Dashboard Overview')
r.bold = True
r.font.size = Pt(24)

doc.add_paragraph('This dashboard helps reviewers understand tax anomalies, identify high-priority exceptions, and investigate mismatches across ETIMS, customs, and withholding data.')

sections = [
    ('1. Purpose', 'The dashboard is designed to support operational review by highlighting unusual tax records and making them easy to interpret at a glance.'),
    ('2. Summary cards', 'The top cards show total findings, matched records, duplicate invoices, and timing gaps so stakeholders can see the overall health of the review quickly.'),
    ('3. Top risk items', 'This section ranks the most urgent issues first so reviewers know what needs attention before anything else.'),
    ('4. Executive summary', 'This panel gives a short management-friendly overview of the current high-risk items and the key issue profile.'),
    ('5. Reconciliation findings', 'This is the detailed table that lists taxpayer-level mismatches between ETIMS and customs values, with a risk score and variance.'),
    ('6. Duplicate invoice numbers', 'This section highlights invoice numbers that appear more than once and shows which buyer taxpayers are involved.'),
    ('7. Timing gaps', 'This table shows unusually long gaps between invoice dates to flag process delays or irregular transaction patterns.'),
    ('8. Detail panel', 'When a row is selected, the dashboard expands into a detailed record view with the relevant taxpayer and transaction context.'),
    ('9. Filters', 'Users can switch between all records, anomalies only, or matched records to focus the review on the most relevant items.'),
    ('10. Print report', 'The print button supports sharing a quick report in meetings or for management documentation.'),
    ('Overall view', 'In short, the dashboard answers three questions: how much is affected, which issues are most critical, and which records need deeper investigation.')
]

for title, content in sections:
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(14)
    doc.add_paragraph(content)

p = doc.add_paragraph()
r = p.add_run('This file documents the dashboard itself, how it works, and why it is useful for stakeholder review.')
r.italic = True

footer = doc.sections[0].footer
footer_para = footer.paragraphs[0]
footer_para.text = 'KRA Tax Anomaly Dashboard Documentation'

doc.save(output_path)
print(f'Wrote document to {output_path}')
